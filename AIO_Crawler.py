import gc
import os
import time
from datetime import datetime
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from crawler import AIOverviewTracker


st.set_page_config(
    page_title="AI Overview Crawler",
    page_icon="🔎",
    layout="wide",
)

st.title("🔎 AI Overview Crawler")
st.caption("上傳提示詞 Excel → 設定參數 → 執行 Google AI Overview 擷取 → 下載結果 Excel")

os.makedirs("output", exist_ok=True)
os.makedirs("screenshots", exist_ok=True)

if "aio_results" not in st.session_state:
    st.session_state.aio_results = []
if "last_output_name" not in st.session_state:
    st.session_state.last_output_name = None


def make_excel_bytes(df: pd.DataFrame, include_screenshots: bool = False) -> bytes:
    """建立可下載的 Excel；可選擇把本機截圖嵌入 Excel。"""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="AIO結果")

    if not include_screenshots or df.empty or "截圖路徑" not in df.columns:
        buffer.seek(0)
        return buffer.getvalue()

    buffer.seek(0)
    wb = load_workbook(buffer)
    ws = wb["AIO結果"]

    widths = {
        "A": 12, "B": 28, "C": 16, "D": 55, "E": 16,
        "F": 28, "G": 65, "H": 18, "I": 35, "J": 22,
        "K": 12, "L": 20,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    img_col = ws.max_column + 1
    ws.cell(row=1, column=img_col, value="搜尋結果截圖")
    ws.column_dimensions[get_column_letter(img_col)].width = 65

    screenshot_col_idx = list(df.columns).index("截圖路徑") + 1

    for row_idx in range(2, ws.max_row + 1):
        path = ws.cell(row=row_idx, column=screenshot_col_idx).value
        if not path or not os.path.exists(path):
            continue
        try:
            img = XLImage(path)
            img.width = 420
            img.height = 260
            ws.add_image(img, ws.cell(row=row_idx, column=img_col).coordinate)
            ws.row_dimensions[row_idx].height = 200
        except Exception:
            pass

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


with st.sidebar:
    st.header("⚙️ 執行設定")

    total_runs = st.number_input(
        "執行次數",
        min_value=1,
        max_value=100,
        value=10,
        step=1,
    )

    delay = st.number_input(
        "每題基礎等待秒數",
        min_value=1,
        max_value=120,
        value=15,
        step=1,
        help="實際每題等待會是此秒數到 +3 秒之間的隨機值。",
    )

    brand_text = st.text_input(
        "品牌關鍵字",
        value="安麗, Amway",
        help="用半形逗號分隔。",
    )
    brand_keywords = [x.strip() for x in brand_text.split(",") if x.strip()]

    chrome_version_text = st.text_input(
        "Chrome 主版本（選填）",
        value="",
        help="通常留空讓 undetected-chromedriver 自動判斷。若有版本相容問題，再填例如 154。",
    )
    chrome_version_main = int(chrome_version_text) if chrome_version_text.strip().isdigit() else None

    include_screenshots_excel = st.checkbox(
        "下載 Excel 時嵌入截圖",
        value=False,
        help="勾選後 Excel 檔會較大。",
    )


st.subheader("① 上傳提示詞 Excel")
uploaded_file = st.file_uploader("選擇 .xlsx 或 .xls", type=["xlsx", "xls"])

if uploaded_file is None:
    st.info("請先上傳 Excel 檔案。")
    st.stop()

file_bytes = uploaded_file.getvalue()
excel = pd.ExcelFile(BytesIO(file_bytes))

c1, c2 = st.columns(2)
with c1:
    selected_sheet = st.selectbox("工作表", excel.sheet_names)

input_df = pd.read_excel(BytesIO(file_bytes), sheet_name=selected_sheet)

with c2:
    keyword_column = st.selectbox(
        "提示詞／關鍵字欄位",
        input_df.columns.tolist(),
        index=0,
    )

keywords = (
    input_df[keyword_column]
    .dropna()
    .astype(str)
    .str.strip()
)
keywords = [x for x in keywords.tolist() if x]

category_map = {}
category_options = ["不使用"] + input_df.columns.tolist()
category_selection = st.selectbox(
    "品類欄位（選填）",
    category_options,
    index=0,
    help="若 Excel 本身有『品類』欄，可在這裡選擇；沒有就維持不使用。",
)
if category_selection != "不使用":
    temp = input_df[[keyword_column, category_selection]].dropna(subset=[keyword_column])
    category_map = {
        str(k).strip(): str(v).strip() if pd.notna(v) else "未分類"
        for k, v in zip(temp[keyword_column], temp[category_selection])
    }

st.success(f"已讀取 {len(keywords):,} 個提示詞。")

with st.expander("預覽上傳資料"):
    st.dataframe(input_df, use_container_width=True, height=300)


st.subheader("② 執行資訊")
total_tasks = len(keywords) * int(total_runs)
m1, m2, m3, m4 = st.columns(4)
m1.metric("提示詞數", f"{len(keywords):,}")
m2.metric("執行輪次", f"{int(total_runs):,}")
m3.metric("總搜尋次數", f"{total_tasks:,}")
m4.metric("最低等待時間", f"{total_tasks * float(delay) / 60:.1f} 分鐘")

st.caption("實際時間還會包含 Google 載入、AIO 展開、引用解析、截圖，以及每題額外 0–3 秒隨機等待。")


st.subheader("③ 執行")
start = st.button(
    "▶ 開始執行",
    type="primary",
    use_container_width=True,
    disabled=(len(keywords) == 0),
)

if start:
    st.session_state.aio_results = []

    progress_bar = st.progress(0.0)
    status_box = st.empty()
    current_box = st.empty()
    metrics_box = st.empty()
    preview_box = st.empty()

    all_results = []
    global_completed = 0
    run_started = datetime.now()
    output_stamp = run_started.strftime("%Y%m%d_%H%M%S")
    autosave_path = os.path.join("output", f"AIO_Result_{output_stamp}_autosave.xlsx")

    try:
        for run_id in range(1, int(total_runs) + 1):
            tracker = None

            try:
                status_box.info(f"Round {run_id} / {int(total_runs)}：啟動 Chrome…")

                tracker = AIOverviewTracker(
                    headless=headless,
                    brand_keywords=brand_keywords,
                    run_id=run_id,
                    category_map=category_map,
                    chrome_version_main=chrome_version_main,
                )

                def on_progress(local_idx, local_total, keyword, result):
                    nonlocal_global = None  # 僅避免 closure 中誤改外層計數
                    completed = (run_id - 1) * len(keywords) + local_idx
                    progress = completed / total_tasks if total_tasks else 1.0

                    current_box.markdown(
                        f"### 🔍 正在處理\n"
                        f"**{keyword}**\n\n"
                        f"Round **{run_id}/{int(total_runs)}** ・ "
                        f"本輪 **{local_idx}/{local_total}**"
                    )
                    progress_bar.progress(min(progress, 1.0))

                    with metrics_box.container():
                        a, b, c, d = st.columns(4)
                        a.metric("整體進度", f"{progress:.1%}")
                        b.metric("已完成", f"{completed}/{total_tasks}")
                        c.metric("AIO", result.get("是否出現AI摘要", ""))
                        c4 = result.get("品牌字內容", "")
                        d.metric("品牌命中", c4 if c4 else "—")

                    # tracker.results 已經包含剛完成的 result。
                    current_rows = all_results + [
                        {
                            **r,
                            "執行輪次": run_id,
                            "執行時間戳": output_stamp,
                        }
                        for r in tracker.results
                    ]
                    preview_box.dataframe(
                        pd.DataFrame(current_rows).tail(20),
                        use_container_width=True,
                        height=420,
                    )

                    # 每題完成就做一次本機 autosave，避免長時間執行中斷後全部遺失。
                    try:
                        pd.DataFrame(current_rows).to_excel(autosave_path, index=False)
                    except Exception:
                        pass

                tracker.check_keywords(
                    keywords,
                    delay=float(delay),
                    progress_callback=on_progress,
                )

                for r in tracker.results:
                    row = dict(r)
                    row["執行輪次"] = run_id
                    row["執行時間戳"] = output_stamp
                    all_results.append(row)

                st.session_state.aio_results = all_results

                # 每輪結束再正式保存一次。
                pd.DataFrame(all_results).to_excel(autosave_path, index=False)

            except Exception as e:
                st.error(f"第 {run_id} 輪執行失敗：{e}")

            finally:
                if tracker is not None:
                    try:
                        tracker.close()
                    except Exception:
                        pass
                gc.collect()
                AIOverviewTracker._kill_residual_chrome()

            if run_id < int(total_runs):
                time.sleep(3)

        st.session_state.aio_results = all_results
        progress_bar.progress(1.0)
        status_box.success(f"🎉 全部完成，共取得 {len(all_results):,} 筆結果。")
        current_box.empty()
        st.session_state.last_output_name = f"AIO_Result_{output_stamp}.xlsx"

    except Exception as e:
        st.session_state.aio_results = all_results
        st.error("執行過程發生未預期錯誤。已完成的資料仍會保留在下方供下載。")
        st.exception(e)


if st.session_state.aio_results:
    st.divider()
    st.subheader("④ 結果與下載")

    result_df = pd.DataFrame(st.session_state.aio_results)
    st.dataframe(result_df, use_container_width=True, height=500)

    excel_bytes = make_excel_bytes(
        result_df,
        include_screenshots=include_screenshots_excel,
    )

    filename = st.session_state.last_output_name or "AIO_Result.xlsx"

    st.download_button(
        "📥 下載完整結果 Excel",
        data=excel_bytes,
        file_name=filename,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )
