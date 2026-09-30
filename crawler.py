import os
import re
import time
import uuid
import random
import tempfile
import shutil
from datetime import datetime
from urllib.parse import urlparse

import requests

import pandas as pd
import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from PIL import Image, ImageDraw
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment

# 檢查是否有 psutil 套件用於清理殘留程序
try:
    import psutil
    HAS_PSUTIL = True
except ImportError:
    HAS_PSUTIL = False

# 預設的品類對照字典（若外部無傳入則以此為基礎）
category_dict = {}

# ✅ debug 開關：True 時會印出驗證失敗的詳細原因，方便排查「有 AI 標識卻抓不到內容」的問題
DEBUG_VERIFY = False


# ============================================================
# AI Overview 追蹤器
# ============================================================
class AIOverviewTracker:
    def __init__(self, headless=False, brand_keywords=None, run_id=None, category_map=None, chrome_version_main=None):
        self.headless = headless
        self.results = []
        self.driver = None
        self._user_data_dir = None
        self.category_dict = category_map or category_dict or {}

        # ============================================================
        # Google /goto 引用網址解析
        # Google 目前可能將 AI Overview 引用網址包成
        # https://www.google.com/goto?...
        # 先用 HTTP redirect 解析，失敗再用 Selenium fallback。
        # cache 可避免同一個 /goto URL 在大量 Prompt 中重複解析。
        # ============================================================
        self.chrome_version_main = chrome_version_main
        self.url_resolve_cache = {}
        self.http_session = requests.Session()
        self.http_session.headers.update({
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/152.0.0.0 Safari/537.36"
            ),
            "Accept-Language": "zh-TW,zh;q=0.9,en;q=0.8",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        })

        # 截圖資料夾
        if run_id:
            self.screenshot_dir = f"screenshots/第{run_id}次_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        else:
            self.screenshot_dir = f"screenshots/{datetime.now().strftime('%Y%m%d_%H%M%S')}"

        os.makedirs(self.screenshot_dir, exist_ok=True)

        self.brand_keywords = brand_keywords or [
            "安麗", "Amway", "紐崔萊", "Nutrilite",
            "雅芝", "Artistry", "益之源", "eSpring",
            "Double X"
        ]

    # ============================================
    # Driver 啟動與健康管理
    # ============================================

    def setup_driver(self, max_retry=3):
        """使用 undetected-chromedriver 啟動（含重試機制 + 自動同意 cookie）"""
        last_error = None

        for attempt in range(1, max_retry + 1):
            driver = None
            try:
                options = uc.ChromeOptions()
                options.add_argument("--no-sandbox")
                options.add_argument("--disable-dev-shm-usage")
                options.add_argument("--disable-gpu")
                options.add_argument("--disable-blink-features=AutomationControlled")
                options.add_argument("--lang=zh-TW")
                options.add_argument("--start-maximized")

                user_dir = os.path.join(
                    tempfile.gettempdir(),
                    f"uc_profile_{uuid.uuid4().hex[:8]}"
                )
                options.add_argument(f"--user-data-dir={user_dir}")

                if self.headless:
                    options.add_argument("--headless=new")

                chrome_kwargs = {"options": options, "use_subprocess": True}
                if self.chrome_version_main:
                    chrome_kwargs["version_main"] = int(self.chrome_version_main)
                driver = uc.Chrome(**chrome_kwargs)

                time.sleep(3)
                driver.set_window_size(1920, 1080)

                # ============================================
                # ✅ 注入 Google Consent Cookie，避免跳出同意頁
                # ============================================
                try:
                    driver.get("https://www.google.com")
                    time.sleep(1.5)

                    driver.add_cookie({
                        "name": "CONSENT",
                        "value": "YES+cb.20220419-08-p0.en+FX+410",
                        "domain": ".google.com",
                        "path": "/",
                    })
                    driver.add_cookie({
                        "name": "SOCS",
                        "value": "CAESHAgBEhJnd3NfMjAyMzA4MTAtMF9SQzIaAmVuIAEaBgiA_LyaBg",
                        "domain": ".google.com",
                        "path": "/",
                    })

                    print("    🍪 已注入 Google Consent Cookie")

                    # 重新載入，使 Cookie 立即生效
                    driver.get("https://www.google.com")
                    time.sleep(1)

                except Exception as e:
                    print(f"    ⚠️ 注入 consent cookie 失敗（不影響啟動）：{e}")
                # ============================================

                driver.get("about:blank")
                _ = driver.current_url

                print(f"✅ Chrome 啟動成功（第 {attempt} 次嘗試）")
                self.driver = driver
                self._user_data_dir = user_dir
                return driver

            except Exception as e:
                last_error = e
                print(f"⚠️ 第 {attempt} 次啟動失敗：{e}")
                try:
                    if driver:
                        driver.quit()
                except:
                    pass
                time.sleep(5)

        raise RuntimeError(f"Driver 啟動失敗（已重試 {max_retry} 次）：{last_error}")

    def _is_driver_alive(self):
        """檢查 driver 是否還活著"""
        try:
            _ = self.driver.current_url
            _ = self.driver.window_handles
            return True
        except Exception:
            return False

    def _restart_driver(self):
        """重啟 driver"""
        print("    🔄 正在重啟 driver...")
        try:
            self.driver.quit()
        except:
            pass
        try:
            if self._user_data_dir:
                shutil.rmtree(self._user_data_dir, ignore_errors=True)
        except:
            pass
        time.sleep(3)
        self._kill_residual_chrome()
        time.sleep(2)
        return self.setup_driver()

    @staticmethod
    def _kill_residual_chrome():
        """清理殘留的 chromedriver 程序"""
        if not HAS_PSUTIL:
            return
        try:
            for proc in psutil.process_iter(['name', 'cmdline']):
                try:
                    name = (proc.info.get('name') or '').lower()
                    cmdline = ' '.join(proc.info.get('cmdline') or []).lower()
                    if 'chromedriver' in name:
                        proc.kill()
                    elif 'chrome' in name and 'uc_profile_' in cmdline:
                        proc.kill()
                except:
                    pass
        except:
            pass

    # ============================================
    # 展開功能
    # ============================================

    def expand_ai_overview_only(self, container):
        """只展開 AI 摘要相關的按鈕（摘要本身的「顯示更多」+ 引用來源的「顯示全部」）

        ✅ 修改：優先在傳入的 container 本身找按鈕，找不到才往上層 ancestor 擴大範圍。
        這樣可以避免容器往上擴大後，誤點到頁面其他區塊（例如自然搜尋結果的評論展開）
        裡同樣叫「顯示更多」的按鈕。
        """
        clicks = 0

        search_scopes = [container]

        try:
            try:
                aio_container = container.find_element(
                    By.XPATH,
                    "ancestor::div[contains(@class, 'kp-') or contains(@data-attrid, 'wa:')]"
                )
                search_scopes.append(aio_container)
            except:
                try:
                    aio_container = container.find_element(By.XPATH, "ancestor::div[5]")
                    search_scopes.append(aio_container)
                except:
                    pass

            # 摘要本身「顯示更多」
            for scope in search_scopes:
                if self._click_button_by_text(scope, "顯示更多"):
                    clicks += 1
                    print("    🔽 已點擊「顯示更多」展開摘要內容")
                    time.sleep(1.2)
                    break

            # 引用來源「顯示全部」
            for scope in search_scopes:
                if self._click_button_by_text(scope, "顯示全部"):
                    clicks += 1
                    print("    🔽 已點擊「顯示全部」展開引用來源")
                    time.sleep(1.2)
                    break

            # 英文版備援
            if clicks == 0:
                for scope in search_scopes:
                    if self._click_button_by_text(scope, "Show more"):
                        clicks += 1
                        print("    🔽 已點擊「Show more」")
                        time.sleep(1)
                        break
                for scope in search_scopes:
                    if self._click_button_by_text(scope, "Show all"):
                        clicks += 1
                        print("    🔽 已點擊「Show all」")
                        time.sleep(1)
                        break

            if clicks == 0:
                print("    ℹ️ 未找到展開按鈕（可能已全部顯示）")
            else:
                print(f"    ✅ 共點擊 {clicks} 次展開按鈕")

            return clicks > 0

        except Exception as e:
            print(f"    ⚠️ 展開時發生錯誤：{e}")
            return False

    def _click_button_by_text(self, container, button_text):
        """精確點擊指定文字的按鈕"""
        try:
            xpaths = [
                f".//span[normalize-space(text())='{button_text}']",
                f".//div[normalize-space(text())='{button_text}']",
                f".//button[normalize-space(text())='{button_text}']",
                f".//a[normalize-space(text())='{button_text}']",
                f".//*[@role='button' and normalize-space(text())='{button_text}']",
                f".//span[contains(text(), '{button_text}')]",
                f".//div[contains(text(), '{button_text}')]",
            ]

            for xpath in xpaths:
                try:
                    elements = container.find_elements(By.XPATH, xpath)
                    for el in elements:
                        if el.is_displayed() and el.is_enabled():
                            href = el.get_attribute("href")
                            if href and "google.com/search" in href:
                                continue

                            el_text = el.text.strip()
                            if len(el_text) > 20:
                                continue

                            try:
                                el.click()
                                return True
                            except:
                                try:
                                    self.driver.execute_script("arguments[0].click();", el)
                                    return True
                                except:
                                    continue
                except:
                    continue

            return False

        except:
            return False

    # ============================================
    # 來源提取功能
    # ============================================


    def _is_google_goto_url(self, url):
        """判斷是否為 Google AI Overview citation 的 /goto 跳轉網址，支援相對與絕對 URL。"""
        try:
            if not url or not isinstance(url, str):
                return False

            url = url.strip()

            if url.startswith("/goto"):
                return True

            parsed = urlparse(url)
            hostname = (parsed.hostname or "").lower()
            path = (parsed.path or "").lower().rstrip("/")

            return (
                path == "/goto"
                and (
                    hostname == "google.com"
                    or hostname.endswith(".google.com")
                )
            )
        except:
            return False

    def _is_excluded_google_url(self, url):
        """排除 Google 本身的搜尋/導覽/帳號等網址。"""
        try:
            parsed = urlparse(url)
            hostname = (parsed.hostname or "").lower()
            path = (parsed.path or "").lower()

            # Google /goto 在解析前先另外處理，因此這裡直接排除。
            if self._is_google_goto_url(url):
                return True

            excluded_hosts_and_paths = [
                ("google.com", "/search"),
                ("google.com", "/url"),
                ("google.com", "/preferences"),
                ("google.com", "/intl"),
                ("google.com", "/webhp"),
                ("google.com", "/advanced_search"),
                ("accounts.google.com", ""),
                ("support.google.com", ""),
                ("policies.google.com", ""),
                ("maps.google.com", ""),
                ("translate.google.com", ""),
                ("play.google.com", ""),
                ("chrome.google.com", ""),
                ("consent.google.com", ""),
            ]

            for excluded_host, excluded_path in excluded_hosts_and_paths:
                if hostname == excluded_host or hostname.endswith("." + excluded_host):
                    if not excluded_path or path.startswith(excluded_path):
                        return True

            return False
        except:
            return False

    def _is_valid_resolved_url(self, url):
        """確認解析後是否得到真正的外部引用網址。"""
        try:
            if not url or not isinstance(url, str):
                return False

            url = url.strip()
            parsed = urlparse(url)

            if parsed.scheme not in ("http", "https"):
                return False

            if not parsed.netloc:
                return False

            if self._is_google_goto_url(url):
                return False

            if self._is_excluded_google_url(url):
                return False

            return True
        except:
            return False


    def _resolve_google_goto(self, goto_url, use_selenium_fallback=True):
        """
        將 Google AI Overview 的 /goto citation URL 解析為實際目的網址。
        支援 href="/goto?..." 與 https://www.google.com/goto?... 兩種形式。
        """
        if not goto_url:
            return None

        goto_url = str(goto_url).strip()

        if goto_url.startswith("/goto"):
            goto_url = "https://www.google.com" + goto_url

        if not self._is_google_goto_url(goto_url):
            return goto_url if self._is_valid_resolved_url(goto_url) else None

        if goto_url in self.url_resolve_cache:
            cached = self.url_resolve_cache[goto_url]
            if cached and DEBUG_VERIFY:
                print(f"      🔁 使用 URL cache：{cached}")
            return cached

        if DEBUG_VERIFY:
            print(f"      🔗 解析 Google /goto：{goto_url[:180]}")

        try:
            response = self.http_session.get(
                goto_url,
                allow_redirects=True,
                timeout=8,
                stream=True,
            )
            final_url = (response.url or "").strip()

            try:
                response.close()
            except:
                pass

            if self._is_valid_resolved_url(final_url):
                self.url_resolve_cache[goto_url] = final_url
                if DEBUG_VERIFY:
                    print(f"      ✅ HTTP 已解析：{final_url}")
                return final_url

        except Exception as e:
            if DEBUG_VERIFY:
                print(f"      ⚠️ HTTP /goto 解析失敗：{e}")

        if not use_selenium_fallback or not self._is_driver_alive():
            self.url_resolve_cache[goto_url] = None
            return None

        driver = self.driver
        original_handle = None
        new_handle = None

        try:
            original_handle = driver.current_window_handle

            driver.switch_to.new_window("tab")
            new_handle = driver.current_window_handle
            driver.get(goto_url)

            deadline = time.time() + 8
            final_url = ""

            while time.time() < deadline:
                final_url = (driver.current_url or "").strip()

                if self._is_valid_resolved_url(final_url):
                    break

                time.sleep(0.25)

            if self._is_valid_resolved_url(final_url):
                self.url_resolve_cache[goto_url] = final_url

                if DEBUG_VERIFY:
                    print(f"      ✅ Selenium 已解析：{final_url}")

                return final_url

        except Exception as e:
            if DEBUG_VERIFY:
                print(f"      ⚠️ Selenium /goto 解析失敗：{e}")

        finally:
            try:
                if new_handle and new_handle in driver.window_handles:
                    driver.switch_to.window(new_handle)
                    driver.close()
            except:
                pass

            try:
                if original_handle and original_handle in driver.window_handles:
                    driver.switch_to.window(original_handle)
                elif driver.window_handles:
                    driver.switch_to.window(driver.window_handles[0])
            except:
                pass

        self.url_resolve_cache[goto_url] = None
        return None


    def _get_aio_citation_links(self, container):
        """
        僅取得 AI Overview citation links。
        不掃描所有 <a>、不使用 .yuRUbf，也不把自然搜尋結果當成 citation。
        """
        links = []
        seen = set()

        selectors = [
            ".//a[contains(@href, '/goto')]",
            ".//a[contains(@href, 'google.com/goto')]",
            ".//div[contains(@class, 'VLkRKc')]//a[@href]",
            ".//div[contains(@class, 'cLjAic')]//a[@href]",
        ]

        for selector in selectors:
            try:
                for link in container.find_elements(By.XPATH, selector):
                    try:
                        href = (link.get_attribute("href") or "").strip()
                    except:
                        href = ""

                    if not href:
                        continue

                    try:
                        organic_ancestors = link.find_elements(
                            By.XPATH,
                            "ancestor::div[contains(@class,'yuRUbf')]"
                        )
                        if organic_ancestors:
                            continue
                    except:
                        pass

                    if href in seen:
                        continue

                    seen.add(href)
                    links.append(link)

            except:
                continue

        return links

    def _find_citation_card(self, link):
        """
        尋找 citation link 所屬的單一來源卡片。
        優先選只包含 1 個 citation link 的最小 parent，
        避免整個 carousel 被視為單一來源。
        """
        candidates = []

        for level in range(1, 7):
            try:
                node = link.find_element(By.XPATH, "/.." * level)
            except:
                continue

            try:
                tag = (node.tag_name or "").lower()
            except:
                tag = ""

            if tag == "g-scrolling-carousel":
                continue

            try:
                text = self.driver.execute_script(
                    "return (arguments[0].innerText || '').trim();", node
                )
            except:
                try:
                    text = node.text.strip()
                except:
                    text = ""

            if not text or len(text) > 1200:
                continue

            try:
                citation_count = len(node.find_elements(
                    By.XPATH,
                    ".//a[contains(@href,'/goto') or contains(@href,'google.com/goto')]"
                ))
            except:
                citation_count = 0

            if citation_count == 1:
                candidates.append((len(text), level, node))

        if candidates:
            candidates.sort(key=lambda x: (x[0], x[1]))
            return candidates[0][2]

        return link

    def _extract_citation_info(self, link):
        """從單一 AIO citation link 擷取來源文字與真正 URL。"""
        try:
            raw_href = (link.get_attribute("href") or "").strip()
            if not raw_href:
                return None

            if raw_href.startswith("/goto"):
                raw_href = "https://www.google.com" + raw_href

            if self._is_google_goto_url(raw_href):
                final_url = self._resolve_google_goto(raw_href)
            else:
                final_url = raw_href

            if not final_url or not self._is_valid_resolved_url(final_url):
                return None

            card = self._find_citation_card(link)

            try:
                card_text = self.driver.execute_script(
                    "return (arguments[0].innerText || '').trim();", card
                )
            except:
                try:
                    card_text = card.text.strip()
                except:
                    card_text = ""

            lines = []
            for line in card_text.splitlines():
                line = " ".join(line.split()).strip()

                if not line:
                    continue

                if line in ("顯示更多", "查看更多", "More", "Show more"):
                    continue

                if line not in lines:
                    lines.append(line)

            source_text = "\n".join(lines).strip()

            if not source_text:
                source_text = (
                    (link.get_attribute("aria-label") or "").strip()
                    or (link.get_attribute("title") or "").strip()
                    or urlparse(final_url).netloc.replace("www.", "")
                )

            if len(source_text) > 1000:
                source_text = source_text[:1000].rstrip() + "..."

            return {
                "標題": source_text,
                "網址": final_url
            }

        except Exception as e:
            if DEBUG_VERIFY:
                print(f"      ⚠️ citation 擷取失敗：{e}")

            return None

    def extract_sources_with_urls(self, container, max_sources=15):
        """
        僅擷取 AI Overview 自己的 citation cards。
        不再掃描所有 a、.yuRUbf、整個容器 cite，也不從摘要文字猜來源。
        """
        sources = []

        try:
            self.expand_ai_overview_only(container)
            time.sleep(0.5)

            links = self._get_aio_citation_links(container)

            for link in links:
                source = self._extract_citation_info(link)

                if not source:
                    continue

                if any(s["網址"] == source["網址"] for s in sources):
                    continue

                sources.append(source)

                if len(sources) >= max_sources:
                    break

            if DEBUG_VERIFY:
                print(f"    🔍 AIO citation links={len(links)}，有效引用={len(sources)}")

        except Exception as e:
            print(f"    ⚠️ 提取 AIO 引用來源時發生錯誤：{e}")

        return sources[:max_sources]

    def _extract_link_info(self, link):
        """從連結元素提取資訊，並將 Google /goto 解析為真正目的網址。"""
        try:
            if not link.is_displayed():
                return None

            href = link.get_attribute("href")

            if not href or not href.startswith('http'):
                return None

            original_href = href

            # ⭐⭐⭐ 核心修正：Google 現在把 AIO citation 包成 /goto
            if self._is_google_goto_url(href):
                href = self._resolve_google_goto(href)

                if not href:
                    if DEBUG_VERIFY:
                        print(f"      ⚠️ 無法解析 Google /goto：{original_href[:180]}")
                    return None

            # 解析後再做 Google /搜尋等排除
            if not self._is_valid_resolved_url(href):
                return None

            # 優先使用 link 本身文字
            text = link.text.strip()

            if not text or len(text) < 2:
                text = link.get_attribute("aria-label") or link.get_attribute("title") or ""

            if not text or len(text) < 2:
                try:
                    parent = link.find_element(By.XPATH, "..")
                    text = parent.text.strip()
                    if len(text) > 80:
                        text = text[:80] + "..."
                except:
                    pass

            if not text or len(text) < 2:
                text = urlparse(href).netloc.replace('www.', '')

            text = ' '.join(text.split())
            if len(text) > 150:
                text = text[:150] + "..."

            return {
                '標題': text,
                '網址': href
            }

        except Exception as e:
            if DEBUG_VERIFY:
                print(f"      ⚠️ _extract_link_info 失敗：{e}")
            return None

    # ============================================
    # AI 摘要偵測功能
    # ============================================



    def _get_ai_summary_text(self, driver, element):
        """
        v5：取得 AI Overview 正文，不再用「刪 DOM parent」的方式。

        核心策略：
        1. 直接讀取 AIO container 的完整 innerText。
        2. 使用 v3 原本已能正確辨識 citation 的：
           _get_aio_citation_links() + _find_citation_card()
           取得引用卡片文字。
        3. 從完整文字中精確扣除 citation card 的文字區塊/行。
        4. 遇到 Google 回饋、匯出等 footer UI 起點後直接停止，
           避免「隱私權政策 / Drive / Gmail / 轉錄中」混入正文。

        注意：
        - 不修改真實 DOM。
        - 不刪 citation parent。
        - 不修改引用資料的擷取邏輯。
        """
        if element is None:
            return ""

        # 1) 先讀取完整 AIO container 文字
        try:
            raw_text = driver.execute_script(
                "return (arguments[0].innerText || '').trim();",
                element
            ) or ""
        except:
            try:
                raw_text = element.text or ""
            except:
                raw_text = ""

        if not raw_text:
            return ""

        def normalize_line(s):
            return " ".join((s or "").replace("\u00a0", " ").split()).strip()

        def normalize_block(s):
            return "\n".join(
                normalize_line(line)
                for line in (s or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
                if normalize_line(line)
            )

        normalized_text = normalize_block(raw_text)
        if not normalized_text:
            return ""

        # 2) 取得「v3 已正確識別」的 citation card 文字
        citation_blocks = []
        citation_lines = set()

        try:
            citation_links = self._get_aio_citation_links(element)
        except:
            citation_links = []

        for link in citation_links:
            try:
                card = self._find_citation_card(link)

                try:
                    card_text = driver.execute_script(
                        "return (arguments[0].innerText || '').trim();",
                        card
                    ) or ""
                except:
                    try:
                        card_text = card.text or ""
                    except:
                        card_text = ""

                block = normalize_block(card_text)
                if not block:
                    continue

                if block not in citation_blocks:
                    citation_blocks.append(block)

                for line in block.split("\n"):
                    line = normalize_line(line)
                    if len(line) >= 2:
                        citation_lines.add(line)

            except:
                continue

        # 3) 優先移除完整 citation block
        working = normalized_text

        # 長的 block 先移除，避免短 block 破壞長 block 的匹配
        for block in sorted(citation_blocks, key=len, reverse=True):
            if len(block) >= 2:
                working = working.replace(block, "\n")

        # 4) 逐行清理：
        #    - 只刪「整行等於 citation card 某一行」的內容
        #    - 不做 substring 刪除，避免正文中提到來源名稱時誤刪整段
        lines = [
            normalize_line(line)
            for line in working.split("\n")
            if normalize_line(line)
        ]

        # Google AIO footer / modal 起點：
        # 一旦出現，後面內容都是互動 UI，不屬於 AI 回答。
        footer_starts = [
            "你的意見能幫助 Google 提升服務品質",
            "你的意見可以幫助 Google 提升服務品質",
            "關於這則回覆",
            "About this response",
        ]

        exact_ui_lines = {
            "謝謝",
            "Thank you",
            "Thanks",
            "顯示更多",
            "顯示全部",
            "查看更多",
            "Show more",
            "Show all",
            "展開",
            "收合",
            "複製",
            "分享",
            "匯出",
            "Export",
            "轉錄中...",
            "轉錄中…",
            "轉錄中",
            "Transcribing...",
            "Transcribing",
        }

        ui_fragments = [
            "詳情請參閱我們的《隱私權政策》",
            "詳情請參閱我們的隱私權政策",
            "儲存至 Google 雲端硬碟",
            "儲存至 Google Drive",
            "儲存至 Gmail",
            "匯出時，你將允許 Google 搜尋",
            "你將允許 Google 搜尋將 AI 生成的資訊",
        ]

        output = []
        seen = set()

        for line in lines:
            # footer 開始後直接停止，不讓後續 UI 混入正文
            if any(line.startswith(marker) for marker in footer_starts):
                break

            # citation card 內容：僅精確整行比對
            if line in citation_lines:
                continue

            if line in exact_ui_lines:
                continue

            if any(fragment in line for fragment in ui_fragments):
                continue

            # 去重
            if line in seen:
                continue

            seen.add(line)
            output.append(line)

        result = "\n".join(output).strip()

        if DEBUG_VERIFY:
            print(
                f"    🔍 [v5 summary] raw={len(normalized_text)}字, "
                f"citation_blocks={len(citation_blocks)}, "
                f"citation_lines={len(citation_lines)}, "
                f"summary={len(result)}字"
            )

        return result

    def _get_full_text(self, driver, element):
        """
        相容介面。
        AIO 文字擷取固定使用 _get_ai_summary_text，
        避免直接讀 container.innerText 而把 citation cards 一起抓進來。
        """
        text = self._get_ai_summary_text(driver, element)

        if text:
            return text.strip()

        try:
            return element.text.strip()
        except:
            return ""


    def _expand_then_get_text(self, driver, container, label=""):
        """先展開 AI Overview，再只讀取生成答案正文。"""
        try:
            self.expand_ai_overview_only(container)
            time.sleep(0.6)
        except:
            pass

        text = self._get_ai_summary_text(driver, container)

        if DEBUG_VERIFY:
            print(f"    🔍 [{label}] 摘要正文長度（已排除 citations）：{len(text)}")

        return text

    def extract_ai_overview_content(self, driver):
        """多重方式偵測 AI 摘要"""
        try:
            driver.execute_script("window.scrollTo(0, 500);")
            time.sleep(0.5)
            driver.execute_script("window.scrollTo(0, 0);")
            time.sleep(0.5)
        except:
            pass

        page_source = driver.page_source
        ai_markers = ["AI 摘要", "AI摘要", "AI Overview", "AI 概覽", "AI概覽"]
        has_ai_marker = any(marker in page_source for marker in ai_markers)

        if not has_ai_marker:
            print("    ℹ️ 頁面未包含 AI 摘要標識")
            return False, "", None, []

        ai_keywords = [
            "AI 摘要", "AI摘要", "AI Overview", "AI 概覽", "AI概覽",
        ]

        for keyword in ai_keywords:
            result = self._find_ai_overview_by_text(driver, keyword)
            if result[0]:
                print(f"    ✅ 透過關鍵字「{keyword}」偵測到 AI 摘要")
                return result

        result = self._find_ai_overview_by_structure(driver)
        if result[0]:
            if self._verify_ai_content(result[1]):
                print("    ✅ 透過 HTML 結構偵測到 AI 摘要")
                return result
            else:
                print("    ⚠️ HTML 結構匹配但內容驗證失敗")

        result = self._find_ai_overview_by_attributes(driver)
        if result[0]:
            if self._verify_ai_content(result[1]):
                print("    ✅ 透過屬性偵測到 AI 摘要")
                return result
            else:
                print("    ⚠️ 屬性匹配但內容驗證失敗")

        result = self._find_ai_overview_by_content(driver)
        if result[0]:
            if self._verify_ai_content(result[1]):
                print("    ✅ 透過內容特徵偵測到 AI 摘要")
                return result
            else:
                print("    ⚠️ 內容特徵匹配但驗證失敗")

        # ✅ 保底機制：頁面明確有 AI 標識，但四種偵測法都沒能通過驗證。
        # 這種情況通常代表容器抓錯（太大/太小）或驗證條件過嚴。
        # 直接對「整個 body」做一次寬鬆的最終嘗試，寧可多抓一點雜訊，也不要漏掉真正的摘要。
        fallback_result = self._fallback_find_ai_overview(driver)
        if fallback_result[0]:
            print("    ✅ 透過保底機制偵測到 AI 摘要")
            return fallback_result

        print("    ℹ️ 頁面有 AI 標識但無法提取有效內容")
        return False, "", None, []

    def _fallback_find_ai_overview(self, driver):
        """
        保底機制：當前面四種偵測法都失敗，但頁面確實有 AI 摘要標識時使用。
        邏輯：以 AI 標識關鍵字所在元素為錠點，往上找到「文字量適中、連結數不過多」的
        容器，只做寬鬆驗證（長度足夠即可），不再套用嚴格的措辭比對。
        """
        try:
            ai_keywords = ["AI 摘要", "AI摘要", "AI Overview", "AI 概覽", "AI概覽"]
            for keyword in ai_keywords:
                try:
                    elements = driver.find_elements(
                        By.XPATH, f"//*[contains(text(),'{keyword}')]"
                    )
                except:
                    continue

                for el in elements:
                    for level in [4, 5, 6, 3, 7, 2, 8, 9]:
                        try:
                            container = el.find_element(By.XPATH, f"ancestor::div[{level}]")
                        except:
                            continue

                        # ✅ 先展開「顯示更多」，再用 textContent/innerText 讀取完整文字
                        # （不只是 .text 可見文字），避免摘要還沒展開就被誤判為內容不足
                        text = self._expand_then_get_text(driver, container, label=f"保底/level{level}")

                        if len(text) < 80:
                            if DEBUG_VERIFY:
                                print(f"    🔍 [保底機制] level={level} 文字過短（{len(text)}字），跳過")
                            continue

                        links = container.find_elements(By.TAG_NAME, "a")
                        if len(links) > 25:
                            # 容器抓太大（可能整個 #rso 或整頁），跳過
                            if DEBUG_VERIFY:
                                print(f"    🔍 [保底機制] level={level} 連結數過多（{len(links)}），跳過")
                            continue

                        clean_text = self._clean_ai_text(text, keyword)
                        if len(clean_text) < 60:
                            if DEBUG_VERIFY:
                                print(f"    🔍 [保底機制] level={level} 清理後過短（{len(clean_text)}字），跳過")
                            continue

                        if DEBUG_VERIFY:
                            print(f"    🔍 [保底機制] ✅ 採用 ancestor::div[{level}]，"
                                  f"文字長度={len(clean_text)}，連結數={len(links)}")

                        sources = self.extract_sources_with_urls(container)
                        return True, clean_text, el, sources

        except Exception as e:
            print(f"    ⚠️ 保底機制錯誤：{e}")

        return False, "", None, []

    def _looks_like_code(self, text):
        """
        ✅ 保險過濾器：偵測文字是否像是 JS/CSS 原始碼混入（而非真正的 AI 摘要）。

        背景：曾發生過抓到的「AI摘要內容」其實是 <script> 標籤裡的
        JS 原始碼（例如 Google Closure Library 的授權宣告開頭）。
        這裡用幾個常見的程式碼特徵做保底攔截，避免同類問題再次發生。
        """
        if not text:
            return False

        code_indicators = [
            "SPDX-License-Identifier" in text,
            "Copyright The Closure Library" in text,
            re.search(r'\bfunction\s*\(', text) is not None,
            re.search(r'\bvar\s+\w+\s*=', text) is not None,
            "getComputedStyle" in text,
            "aria-hidden" in text and "setAttribute" in text,
            text.count("{") > 5 and text.count(";") > 5,
        ]

        return sum(1 for c in code_indicators if c) >= 2

    def _verify_ai_content(self, text):
        """
        二次驗證內容是否真的是 AI 摘要。

        ✅ 修改說明：原本任何一個排除詞命中、或行數不足、或沒中任何 AI 用語，
        就會「一票否決」整段判定為不是 AI 摘要。這在容器抓取範圍稍有偏差
        （包進旁邊的「查看更多」按鈕，或摘要本身是精簡的單段文字）時，
        會誤殺明明存在的 AI 摘要。

        改為「加分制」：綜合長度、行數、AI 用語命中數等多個訊號打分，
        達到門檻即視為通過；排除詞只在「開頭就是雜訊區塊」時才生效，
        不再對整段全文搜尋。
        """
        if not text:
            if DEBUG_VERIFY:
                print("    🔍 [verify失敗] 文字為空")
            return False

        if self._looks_like_code(text):
            if DEBUG_VERIFY:
                print(f"    🔍 [verify失敗] 內容疑似程式碼混入：{text[:100]!r}")
            return False

        if len(text) < 60:
            if DEBUG_VERIFY:
                print(f"    🔍 [verify失敗] 長度不足：{len(text)} 字，內容：{text[:80]!r}")
            return False

        # 排除詞：只在「文字開頭幾行」明顯是雜訊區塊（相關搜尋/購物/廣告等）時才排除，
        # 避免因為容器範圍稍大、混入旁邊按鈕文字而誤殺真正的摘要
        exclude_patterns = [
            "搜尋結果", "相關搜尋", "其他人也搜尋了", "其他人也問了",
            "People also ask", "熱門搜尋",
            "贊助商", "查看更多", "更多結果",
        ]
        head_text = text[:40]
        for pattern in exclude_patterns:
            if head_text.startswith(pattern) or pattern in head_text:
                if DEBUG_VERIFY:
                    print(f"    🔍 [verify失敗] 開頭疑似雜訊區塊，命中「{pattern}」："
                          f"{text[:80]!r}")
                return False

        lines = [l.strip() for l in text.split('\n') if l.strip()]

        ai_phrases = [
            "以下是", "以下為", "主要包括", "可能包括",
            "通常", "一般來說", "根據", "建議",
            "首先", "其次", "此外", "包含", "例如",
            "常見", "重要的是", "需要注意", "值得注意",
            "總結來說", "簡單來說", "換句話說", "另外",
            "可以", "會", "是指", "指的是",
        ]
        phrase_count = sum(1 for phrase in ai_phrases if phrase in text)

        # 加分制評分：
        # - 長度 >= 150 字：+1
        # - 至少 2 個非空行（多行結構，像列點或分段）：+1
        # - 命中至少 1 個 AI 常見用語：+1
        # - 純長度 >= 300 字（即使沒中用語、只有一行，也很可能是完整摘要段落）：+1
        score = 0
        if len(text) >= 150:
            score += 1
        if len(lines) >= 2:
            score += 1
        if phrase_count >= 1:
            score += 1
        if len(text) >= 300:
            score += 1

        passed = score >= 1  # 只要符合任一訊號即放行，寧可多抓，不要漏抓

        if DEBUG_VERIFY:
            print(f"    🔍 [verify] 長度={len(text)}, 行數={len(lines)}, "
                  f"命中AI用語數={phrase_count}, 分數={score}, "
                  f"結果={'通過' if passed else '未通過'}")
            if not passed:
                print(f"    🔍 [verify未通過內容預覽] {text[:150]!r}")

        return passed

    def _find_ai_overview_by_text(self, driver, keyword):
        """透過文字關鍵字尋找 AI 摘要"""
        try:
            xpaths = [
                f"//*[contains(text(),'{keyword}')]",
                f"//*[contains(@aria-label,'{keyword}')]",
                f"//*[contains(@title,'{keyword}')]",
            ]

            for xpath in xpaths:
                try:
                    elements = driver.find_elements(By.XPATH, xpath)

                    for el in elements:
                        try:
                            if not el.is_displayed():
                                driver.execute_script(
                                    "arguments[0].scrollIntoView({block: 'center'});", el
                                )
                                time.sleep(0.3)
                        except:
                            pass

                        container = self._find_ai_container(el)
                        if not container:
                            continue

                        # ✅ 先展開「顯示更多」再讀取文字，避免抓到截斷前的預覽片段
                        text = self._expand_then_get_text(driver, container, label="by_text")
                        if len(text) < 50:
                            if DEBUG_VERIFY:
                                print(f"    🔍 [by_text/{keyword}] 容器文字過短（{len(text)}字），跳過")
                            continue

                        clean_text = self._clean_ai_text(text, keyword)
                        if len(clean_text) < 30:
                            if DEBUG_VERIFY:
                                print(f"    🔍 [by_text/{keyword}] 清理後文字過短（{len(clean_text)}字），跳過")
                            continue

                        if not self._verify_ai_content(clean_text):
                            continue

                        sources = self.extract_sources_with_urls(container)
                        return True, clean_text, el, sources

                except:
                    continue

        except Exception as e:
            print(f"    ⚠️ 文字偵測錯誤：{e}")

        return False, "", None, []

    def _find_ai_overview_by_structure(self, driver):
        """透過 HTML 結構偵測"""
        try:
            selectors = [
                "[data-attrid='wa:/description']",
                ".Wt5Tfe",
                ".IZ6rdc",
                # ✅ 新增備援選擇器：Google 前端 class 常常改版，
                # 這裡多加幾個近期版本觀察到的容器特徵，降低單一 class 失效時整段偵測不到的風險
                "[data-attrid*='wa:']",
                "div[jsname][data-attrid]",
            ]

            for selector in selectors:
                try:
                    elements = driver.find_elements(By.CSS_SELECTOR, selector)
                    for el in elements:
                        # ✅ 先展開再讀取完整文字，避免抓到摺疊前的短預覽
                        text = self._expand_then_get_text(driver, el, label=f"by_structure/{selector}")
                        if len(text) <= 150:
                            if DEBUG_VERIFY:
                                print(f"    🔍 [by_structure/{selector}] 文字過短（{len(text)}字），跳過")
                            continue
                        if not self._looks_like_ai_content(text):
                            continue
                        links = el.find_elements(By.TAG_NAME, "a")
                        if len(links) > 15:
                            if DEBUG_VERIFY:
                                print(f"    🔍 [by_structure/{selector}] 連結數過多（{len(links)}），跳過")
                            continue
                        clean_text = self._clean_ai_text(text)
                        sources = self.extract_sources_with_urls(el)
                        return True, clean_text, el, sources
                except:
                    continue

        except Exception as e:
            print(f"    ⚠️ 結構偵測錯誤：{e}")

        return False, "", None, []

    def _find_ai_overview_by_attributes(self, driver):
        """透過屬性偵測"""
        try:
            attr_selectors = [
                "[data-attrid*='description']",
                "[data-attrid*='overview']",
                "[aria-label*='AI']",
                "[aria-label*='摘要']",
                "[aria-label*='概覽']",
            ]

            for selector in attr_selectors:
                try:
                    elements = driver.find_elements(By.CSS_SELECTOR, selector)
                    for el in elements:
                        # ✅ 先展開再讀取完整文字
                        text = self._expand_then_get_text(driver, el, label=f"by_attributes/{selector}")
                        if len(text) <= 100:
                            if DEBUG_VERIFY:
                                print(f"    🔍 [by_attributes/{selector}] 文字過短（{len(text)}字），跳過")
                            continue
                        if not self._looks_like_ai_content(text):
                            continue
                        links = el.find_elements(By.TAG_NAME, "a")
                        if len(links) > 15:
                            if DEBUG_VERIFY:
                                print(f"    🔍 [by_attributes/{selector}] 連結數過多（{len(links)}），跳過")
                            continue
                        clean_text = self._clean_ai_text(text)
                        sources = self.extract_sources_with_urls(el)
                        return True, clean_text, el, sources
                except:
                    continue

        except Exception as e:
            print(f"    ⚠️ 屬性偵測錯誤：{e}")

        return False, "", None, []

    def _find_ai_overview_by_content(self, driver):
        """透過內容特徵偵測"""
        try:
            candidates = driver.find_elements(
                By.CSS_SELECTOR,
                "#rso > div:first-child, #rcnt > div:first-child > div:first-child, .kp-wholepage-pane, .ULSxyf"
            )

            for candidate in candidates:
                # ✅ 先展開再讀取完整文字
                text = self._expand_then_get_text(driver, candidate, label="by_content")
                if len(text) <= 200:
                    if DEBUG_VERIFY:
                        print(f"    🔍 [by_content] 文字過短（{len(text)}字），跳過")
                    continue
                if not self._looks_like_ai_content(text):
                    continue
                links = candidate.find_elements(By.TAG_NAME, "a")
                if len(links) >= 15:
                    if DEBUG_VERIFY:
                        print(f"    🔍 [by_content] 連結數過多（{len(links)}），跳過")
                    continue
                clean_text = self._clean_ai_text(text)
                sources = self.extract_sources_with_urls(candidate)
                return True, clean_text, candidate, sources

        except Exception as e:
            print(f"    ⚠️ 內容偵測錯誤：{e}")

        return False, "", None, []


    def _find_ai_container(self, element):
        """
        找 AI Overview 的合理容器。
        優先含 citation、organic result 少、文字範圍較小的 ancestor，
        避免升到整個 #rso / SERP。
        """
        driver = self.driver
        candidates = []

        for level in [1, 2, 3, 4, 5, 6, 7, 8, 9]:
            try:
                container = element.find_element(
                    By.XPATH,
                    f"ancestor::div[{level}]"
                )
            except:
                continue

            try:
                raw_text = driver.execute_script(
                    "return (arguments[0].innerText || '').trim();",
                    container
                )
            except:
                try:
                    raw_text = container.text.strip()
                except:
                    raw_text = ""

            text_len = len(raw_text)

            if text_len < 60:
                continue

            try:
                goto_count = len(container.find_elements(
                    By.XPATH,
                    ".//a[contains(@href,'/goto') or contains(@href,'google.com/goto')]"
                ))
            except:
                goto_count = 0

            try:
                legacy_count = len(container.find_elements(
                    By.CSS_SELECTOR,
                    ".VLkRKc a, .cLjAic a"
                ))
            except:
                legacy_count = 0

            try:
                organic_count = len(
                    container.find_elements(
                        By.CSS_SELECTOR,
                        ".yuRUbf"
                    )
                )
            except:
                organic_count = 0

            citation_count = goto_count + legacy_count

            if organic_count > 2:
                continue

            score = (
                (10000 if citation_count > 0 else 0)
                + min(citation_count, 20) * 100
                - organic_count * 3000
                - min(text_len, 10000) / 10
                - level
            )

            candidates.append(
                (
                    score,
                    citation_count,
                    organic_count,
                    text_len,
                    level,
                    container
                )
            )

        if candidates:
            candidates.sort(
                key=lambda x: x[0],
                reverse=True
            )

            best = candidates[0]

            if DEBUG_VERIFY:
                preview = [
                    {
                        "level": c[4],
                        "text": c[3],
                        "citation": c[1],
                        "organic": c[2],
                        "score": round(c[0], 1),
                    }
                    for c in candidates[:6]
                ]

                print(
                    f"    🔍 [find_container] 候選：{preview}；"
                    f"選 level={best[4]}"
                )

            return best[5]

        try:
            return element.find_element(
                By.XPATH,
                "ancestor::div[1]"
            )
        except:
            return element

    def _looks_like_ai_content(self, text):
        """
        判斷文字是否像 AI 生成的內容。

        ✅ 修改說明：原本用「exclude_indicators 只要中一個就整段排除」的方式，
        容易因為容器裡混入購物/廣告相關的雜訊字樣而誤殺。改為採計分制，
        只有在雜訊訊號明顯偏多、且 AI 訊號不足時才判定為不是 AI 內容。
        """
        if not text or len(text) < 80:
            if DEBUG_VERIFY:
                print(f"    🔍 [looks_like失敗] 長度不足：{len(text) if text else 0}")
            return False

        if self._looks_like_code(text):
            if DEBUG_VERIFY:
                print(f"    🔍 [looks_like失敗] 內容疑似程式碼混入：{text[:100]!r}")
            return False

        head = text[:150]

        exclude_hits = sum([
            "廣告" in head,
            "贊助" in head,
            "搜尋結果" in text[:50],
            "相關搜尋" in text[:50],
            "其他人也問了" in head,
            "People also ask" in head,
            text.count("$") > 5,
            text.count("NT$") > 4,
            "立即購買" in text,
            "加入購物車" in text,
            text.count("http") > 8,
        ])

        ai_signal_count = sum([
            len(text) > 200,
            text.count('\n') >= 2,
            any(phrase in text for phrase in [
                "以下是", "以下為", "主要包括", "可能包括",
                "通常", "一般來說", "根據", "研究顯示",
                "建議", "需要注意", "值得注意",
                "首先", "其次", "最後", "此外",
                "包含", "例如", "比如", "可能是",
                "常見的", "重要的是", "總結來說",
            ]),
        ])

        # 雜訊訊號 >= 3 個才視為明顯不是 AI 內容；否則只要有基本 AI 訊號就放行
        if exclude_hits >= 3:
            if DEBUG_VERIFY:
                print(f"    🔍 [looks_like失敗] 雜訊訊號過多：{exclude_hits} 個，"
                      f"內容預覽：{text[:100]!r}")
            return False

        passed = ai_signal_count >= 1 or len(text) >= 250

        if DEBUG_VERIFY and not passed:
            print(f"    🔍 [looks_like失敗] AI訊號不足：{ai_signal_count}，"
                  f"長度={len(text)}，內容預覽：{text[:100]!r}")

        return passed


    def _clean_ai_text(self, text, keyword_to_remove=None):
        """
        v5：只做低風險的文字清理。
        citation 已在 _get_ai_summary_text() 扣除，
        這裡不再碰來源內容，也不使用模糊 DOM/文字推測。
        """
        if not text:
            return ""

        lines = [
            " ".join(line.replace("\u00a0", " ").split()).strip()
            for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
            if line and line.strip()
        ]

        heading_texts = {
            "AI 摘要",
            "AI摘要",
            "AI Overview",
            "AI 概覽",
            "AI概覽",
            "生成式 AI",
            "生成式AI",
            "Generative AI",
        }

        if keyword_to_remove:
            heading_texts.add(keyword_to_remove.strip())

        exact_ui_lines = {
            "顯示更多", "顯示全部", "Show more", "Show all",
            "展開", "收合", "更多資訊", "查看更多",
            "複製", "分享", "匯出", "Export",
            "謝謝", "Thank you", "Thanks",
            "轉錄中...", "轉錄中…", "轉錄中",
            "Transcribing...", "Transcribing",
        }

        footer_starts = [
            "你的意見能幫助 Google 提升服務品質",
            "你的意見可以幫助 Google 提升服務品質",
            "關於這則回覆",
            "About this response",
        ]

        result = []
        seen = set()

        for line in lines:
            if any(line.startswith(marker) for marker in footer_starts):
                break

            # 僅刪除「完全等於標題」的行，不用 startswith，
            # 避免正文剛好以相似詞開頭而被整行刪除。
            if line in heading_texts:
                continue

            if line in exact_ui_lines:
                continue

            if line in seen:
                continue

            seen.add(line)
            result.append(line)

        return "\n".join(result).strip()

    def check_brand_keywords(self, text):
        """檢查品牌關鍵字"""
        if not text:
            return False, []

        found_brands = []
        text_lower = text.lower()

        for brand in self.brand_keywords:
            if brand.lower() in text_lower:
                found_brands.append(brand)

        has_brand = len(found_brands) > 0
        return has_brand, found_brands

    def extract_sources_from_text(self, text):
        """從文字提取引用來源（備用）"""
        if not text:
            return []

        sources = []
        lines = text.split('\n')

        source_patterns = [
            r'(?:來源|資料來源|引用自|參考資料|根據|according to)[：:]\s*(.+)',
            r'(.+?)(?:指出|表示|說明|報導|顯示|提到)',
            r'(?:維基百科|Wikipedia|百度百科|台灣大百科).*',
            r'(https?://[^\s]+)',
        ]

        for line in lines:
            line = line.strip()
            if not line or len(line) < 5:
                continue

            for pattern in source_patterns:
                match = re.search(pattern, line, re.IGNORECASE)
                if match:
                    try:
                        source = match.group(1) if match.lastindex and match.lastindex >= 1 else match.group(0)
                        source = source.strip()
                        source = re.sub(r'[，。、！？]$', '', source)
                        if source and len(source) > 3 and source not in sources:
                            sources.append(source)
                    except:
                        pass
                    break

        return sources[:10]

    def take_screenshot(self, driver, index, keyword, ai_element=None):
        """截圖並標記 AI 摘要"""
        safe_keyword = re.sub(r'[\\/:*?"|]', "_", keyword)
        path = os.path.join(self.screenshot_dir, f"{index}_{safe_keyword}.png")

        try:
            total_height = driver.execute_script(
                "return Math.max(document.body.scrollHeight, document.documentElement.scrollHeight);"
            )

            original_size = driver.get_window_size()
            driver.set_window_size(original_size["width"], min(total_height, 10000))
            time.sleep(0.5)

            driver.save_screenshot(path)
            driver.set_window_size(original_size["width"], original_size["height"])

            if ai_element:
                try:
                    img = Image.open(path)
                    draw = ImageDraw.Draw(img)

                    loc = ai_element.location
                    size = ai_element.size

                    left = loc["x"]
                    top = loc["y"]
                    right = left + size["width"]
                    bottom = top + size["height"]

                    draw.rectangle([left, top, right, bottom], outline="red", width=6)
                    img.save(path)
                except:
                    pass
        except Exception as e:
            print(f"  ⚠️ 截圖失敗：{e}")
            return ""

        return path

    def _handle_consent_page(self, driver, original_url=None):
        """
        ✅ 保底機制：偵測並自動點擊 Google 同意頁的「全部接受」按鈕
        original_url: 若被重導向至同意頁，接受後要重新開啟的目標網址
        """
        try:
            current_url = driver.current_url
            page_source_lower = driver.page_source.lower()

            is_consent_page = (
                "consent.google.com" in current_url
                or "before you continue" in page_source_lower
                or "在您繼續操作前" in driver.page_source
                or "繼續使用 google 服務前" in driver.page_source
            )

            if not is_consent_page:
                return False

            print("    🍪 偵測到 Google 同意頁，嘗試自動接受...")

            accept_texts = [
                "全部接受", "接受全部", "我同意", "接受",
                "Accept all", "I agree", "Agree"
            ]

            for text in accept_texts:
                try:
                    xpaths = [
                        f"//button[.//*[contains(text(), '{text}')]]",
                        f"//button[contains(., '{text}')]",
                        f"//div[@role='button' and contains(., '{text}')]",
                    ]
                    for xpath in xpaths:
                        btns = driver.find_elements(By.XPATH, xpath)
                        for btn in btns:
                            if btn.is_displayed() and btn.is_enabled():
                                try:
                                    btn.click()
                                except:
                                    driver.execute_script("arguments[0].click();", btn)

                                print(f"    ✅ 已點擊「{text}」，同意頁已自動處理")
                                time.sleep(2)

                                if original_url:
                                    driver.get(original_url)
                                    time.sleep(2)

                                return True
                except:
                    continue

            print("    ⚠️ 找不到自動同意按鈕，可能需要人工確認")
            return False

        except Exception as e:
            print(f"    ⚠️ 處理同意頁時發生錯誤：{e}")
            return False

    def check_keyword(self, driver, keyword, index):
            """檢查單一關鍵字"""
            print(f"\n[{index}] 檢查：{keyword}")
            encountered_captcha = False

            try:
                url = f"https://www.google.com/search?q={keyword}&gl=tw&hl=zh-TW"
                driver.get(url)

                time.sleep(2)

                # ✅ 保底：檢查並處理 Google 同意頁
                self._handle_consent_page(driver, original_url=url)

                # ============================================================
                # ✅ 精確偵測驗證碼頁面 (修正原本 page_source 範圍過寬導致的誤判)
                # ============================================================
                current_url = driver.current_url.lower()
                page_title = driver.title.lower()

                # Google 驗證阻擋頁網址一定會包含 /sorry/
                is_captcha_page = (
                    "google.com/sorry/" in current_url
                    or "unusual traffic" in page_title
                    or "異常流量" in page_title
                )

                if is_captcha_page:
                    print("  ⚠️ 偵測到【真正的驗證阻擋頁面】，等待 300 秒後重試...")
                    encountered_captcha = True
                    time.sleep(300)
                    driver.refresh()
                    time.sleep(2)

                    # 重試後再做一次同意頁與驗證偵測
                    self._handle_consent_page(driver, original_url=url)

                    current_url = driver.current_url.lower()
                    page_title = driver.title.lower()
                    if "google.com/sorry/" in current_url or "unusual traffic" in page_title:
                        print("  ⚠️ 重試後仍在驗證頁面，本次檢查可能無法取得資料")
                # ============================================================

                time.sleep(random.uniform(2, 4))

                driver.execute_script("window.scrollTo(0, 300);")
                time.sleep(random.uniform(0.5, 1))
                driver.execute_script("window.scrollTo(0, 0);")

                WebDriverWait(driver, 10).until(
                    EC.presence_of_element_located((By.TAG_NAME, "body"))
                )

                has_ai, clean_text, ai_element, sources_with_urls = self.extract_ai_overview_content(driver)

                if has_ai and ai_element:
                    try:
                        driver.execute_script(
                            "arguments[0].scrollIntoView({block: 'center'});",
                            ai_element
                        )
                        time.sleep(0.8)
                    except:
                        pass

                screenshot_path = self.take_screenshot(driver, index, keyword, ai_element)

                if has_ai and clean_text:
                    has_brand, found_brands = self.check_brand_keywords(clean_text)
                    brand_status = "是" if has_brand else "否"
                    brand_details = "、".join(found_brands) if found_brands else ""

                    if sources_with_urls:
                        sources_text = "\n".join([
                            f"{i + 1}. {s['標題']}\n   {s['網址']}"
                            for i, s in enumerate(sources_with_urls)
                        ])
                    else:
                        # 沒有在 AIO citation DOM 中找到來源，就保持空白。
                        # 不再從 AI摘要內容 猜測來源，避免製造錯誤引用。
                        sources_text = ""

                    print("  ✓ 偵測到 AI 摘要")
                    if has_brand:
                        print(f"    🏷️ 包含 brand 字：{brand_details}")
                    if sources_with_urls:
                        print(f"    📎 找到 {len(sources_with_urls)} 個引用來源（含網址）")
                else:
                    brand_status = ""
                    brand_details = ""
                    sources_text = ""
                    print("  ✗ 未出現 AI 摘要")

                if encountered_captcha:
                    print("  🔒 本次檢查曾遇到驗證頁面")

                return {
                    "品類": self.get_category(keyword),
                    "關鍵字": keyword,
                    "是否出現AI摘要": "是" if has_ai else "否",
                    "AI摘要內容": clean_text if has_ai else "",
                    "是否出現品牌字": brand_status,
                    "品牌字內容": brand_details,
                    "引用資料": sources_text,
                    "是否出現驗證頁面": "是" if encountered_captcha else "否",
                    "截圖路徑": screenshot_path,
                    "檢查時間": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                }

            except Exception as e:
                print(f"  ❌ 錯誤：{e}")
                return {
                    "品類": self.get_category(keyword),
                    "關鍵字": keyword,
                    "是否出現AI摘要": "錯誤",
                    "AI摘要內容": str(e),
                    "是否出現品牌字": "",
                    "品牌字內容": "",
                    "引用資料": "",
                    "是否出現驗證頁面": "是" if encountered_captcha else "否",
                    "截圖路徑": "",
                    "檢查時間": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                }

    def check_keywords(self, keywords, delay=5, progress_callback=None):
        """批次檢查關鍵字；每完成一題可透過 callback 回報 Streamlit UI。"""
        driver = self.setup_driver()
        total = len(keywords)
        try:
            for idx, kw in enumerate(keywords, 1):
                if not self._is_driver_alive():
                    print("⚠️ 偵測到 driver 已失效，嘗試重啟...")
                    try:
                        driver = self._restart_driver()
                    except Exception as e:
                        print(f"❌ 重啟失敗：{e}")
                        for remain_idx, remain_kw in enumerate(keywords[idx - 1:], idx):
                            result = {
                                "品類": self.get_category(remain_kw), "關鍵字": remain_kw,
                                "是否出現AI摘要": "錯誤", "AI摘要內容": "Driver 失效無法重啟",
                                "是否出現品牌字": "", "品牌字內容": "", "引用資料": "",
                                "是否出現驗證頁面": "否", "截圖路徑": "",
                                "檢查時間": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                            }
                            self.results.append(result)
                            if progress_callback:
                                progress_callback(remain_idx, total, remain_kw, result)
                        break

                try:
                    result = self.check_keyword(driver, kw, idx)
                except Exception as e:
                    result = {
                        "品類": self.get_category(kw), "關鍵字": kw,
                        "是否出現AI摘要": "錯誤", "AI摘要內容": str(e),
                        "是否出現品牌字": "", "品牌字內容": "", "引用資料": "",
                        "是否出現驗證頁面": "否", "截圖路徑": "",
                        "檢查時間": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    }
                self.results.append(result)
                if progress_callback:
                    progress_callback(idx, total, kw, result)

                if idx < total:
                    wait_time = random.uniform(delay, delay + 3)
                    print(f"   ⏳ 等待 {wait_time:.1f} 秒...")
                    time.sleep(wait_time)
        finally:
            self.close()
        return self.results

    def close(self):
        """安全關閉瀏覽器與暫存 profile。"""
        try:
            if self.driver:
                self.driver.quit()
        except Exception:
            pass
        self.driver = None
        try:
            if self._user_data_dir:
                shutil.rmtree(self._user_data_dir, ignore_errors=True)
        except Exception:
            pass
        self._user_data_dir = None

    def get_category(self, keyword):
        """根據關鍵字取得品類"""
        # 1. 精確比對
        if keyword in self.category_dict:
            return self.category_dict[keyword]

        # 2. 模糊比對（互相包含）
        for key, cat in self.category_dict.items():
            if key in keyword or keyword in key:
                return cat

        return "未分類"

    def save_to_excel(self, filename="ai_overview_results.xlsx"):
        """儲存結果到 Excel（含截圖）"""
        df = pd.DataFrame(self.results)

        column_order = [
            "品類", "關鍵字", "是否出現AI摘要", "AI摘要內容",
            "是否出現品牌字", "品牌字內容", "引用資料",
            "是否出現驗證頁面", "截圖路徑", "檢查時間"
        ]

        for col in column_order:
            if col not in df.columns:
                df[col] = ""

        df = df[column_order]
        df.to_excel(filename, index=False)

        wb = load_workbook(filename)
        ws = wb.active

        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 50
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 30
        ws.column_dimensions['G'].width = 60
        ws.column_dimensions['H'].width = 18
        ws.column_dimensions['I'].width = 30
        ws.column_dimensions['J'].width = 20

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')

        img_col = ws.max_column + 1
        ws.cell(row=1, column=img_col, value="搜尋結果截圖")

        screenshot_col_idx = df.columns.get_loc("截圖路徑") + 1

        for row in range(2, ws.max_row + 1):
            path = ws.cell(row=row, column=screenshot_col_idx).value
            if not path or not os.path.exists(path):
                continue

            try:
                img = XLImage(path)
                img.width = 420
                img.height = 260
                ws.add_image(img, ws.cell(row=row, column=img_col).coordinate)
                ws.row_dimensions[row].height = 200
            except Exception as e:
                print(f"  ⚠️ 無法插入截圖 {path}: {e}")

        ws.column_dimensions[get_column_letter(img_col)].width = 65

        wb.save(filename)
        print(f"\n📊 Excel 已儲存：{filename}")

        return df